#!/usr/bin/env python3
"""Generate + upload an ACT Agent recording (script .py + params workbook) to MinIO/S3.

Given a recorded Playwright ``.py`` and a params JSON, this:
  1. builds the sibling params file (``.xlsx`` by default, ``.csv`` optional) in the exact
     layout the runner reads (active sheet ``params``: header row + one data row per set),
  2. uploads both to ``recordings/<name>/<name>.py`` and ``recordings/<name>/<name>_params.<ext>``
     in the storage bucket (``TENANT_ID`` if set, else ``STORAGE_ACTIVITIES_BUCKET``), and
  3. prints the ``aetherion agent 'ACT Agent'`` command to run it.

Params JSON accepts any of:
  - ``{"params": [ {..} ]}``                         (the workbook-style payload)
  - ``[ {..}, {..} ]``                               (list of param sets -> multiple data rows)
  - ``{..}``                                         (a single param set)

Examples:
  scripts/upload_recording.py --py rec.py --params params.json
  scripts/upload_recording.py --py rec.py --params-json '{"params":[{"url":"..."}]}' --name My_Rec_v1.0
  scripts/upload_recording.py --py rec.py --params params.json --format csv --dry-run
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import sys
from pathlib import Path
from typing import Any


# --------------------------------------------------------------------------- payload parsing
def normalize_param_sets(payload: Any) -> list[dict[str, str]]:
    """Return the ordered list of param-set dicts from any accepted payload shape."""
    if isinstance(payload, dict) and "params" in payload:
        raw_sets = payload.get("params") or []
    elif isinstance(payload, list):
        raw_sets = payload
    elif isinstance(payload, dict):
        raw_sets = [payload]
    else:
        raise ValueError("params payload must be a dict, a list of dicts, or {'params': [...]}")

    sets: list[dict[str, str]] = []
    for entry in raw_sets:
        if not isinstance(entry, dict):
            raise ValueError(f"each param set must be an object, got: {type(entry).__name__}")
        sets.append({str(k): _cell(v) for k, v in entry.items()})
    if not sets:
        raise ValueError("no parameter sets found in payload")
    return sets


def _cell(value: Any) -> str:
    """Render a value as the runner reads it (every cell becomes a trimmed string)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (list, tuple)):
        return "|".join(_cell(v) for v in value)
    return str(value)


def _ordered_headers(rows: list[dict[str, str]]) -> list[str]:
    """Union of keys across rows, preserving first-seen order."""
    headers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                headers.append(key)
    return headers


# --------------------------------------------------------------------------- file builders
def build_params_csv(param_sets: list[dict[str, str]]) -> bytes:
    headers = _ordered_headers(param_sets)
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(headers)
    for row in param_sets:
        writer.writerow([row.get(h, "") for h in headers])
    return buf.getvalue().encode("utf-8")


def build_params_xlsx(param_sets: list[dict[str, str]]) -> bytes:
    import openpyxl  # lazy: only needed for xlsx output

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "params"
    headers = _ordered_headers(param_sets)
    ws.append(headers)
    for row in param_sets:
        ws.append([row.get(h, "") for h in headers])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- env / storage
def load_env(env_path: Path) -> dict[str, str]:
    """Parse a dotenv file; real os.environ overrides it for the keys we care about."""
    cfg: dict[str, str] = {}
    if env_path.is_file():
        for line in env_path.read_text().splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            cfg[key.strip()] = value.strip().strip('"').strip("'")
    for key in ("TENANT_ID", "STORAGE_ENDPOINT", "STORAGE_ACCESS_KEY", "STORAGE_SECRET_KEY", "STORAGE_ACTIVITIES_BUCKET"):
        if os.environ.get(key):
            cfg[key] = os.environ[key]
    return cfg


def resolve_bucket(cfg: dict[str, str], override: str | None) -> str:
    # Mirror the runner's _get_bucket_name(): TENANT_ID wins, else STORAGE_ACTIVITIES_BUCKET.
    bucket = (override or cfg.get("TENANT_ID") or cfg.get("STORAGE_ACTIVITIES_BUCKET") or "").strip()
    if not bucket:
        raise SystemExit("ERROR: no bucket: set --bucket, or TENANT_ID / STORAGE_ACTIVITIES_BUCKET in --env.")
    return bucket


def make_s3(cfg: dict[str, str]):
    import boto3  # lazy

    endpoint = cfg.get("STORAGE_ENDPOINT") or None
    return boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=cfg.get("STORAGE_ACCESS_KEY"),
        aws_secret_access_key=cfg.get("STORAGE_SECRET_KEY"),
        region_name=cfg.get("AWS_REGION") or "us-east-1",
    )


def run_command(name: str, file_key: str, execution_mode: str) -> str:
    payload = {
        "test_suite_id": name,
        "recordings": [{"id": name, "name": name, "file": file_key}],
        "execution_mode": execution_mode,
    }
    return "./.venv/bin/aetherion agent 'ACT Agent' '" + json.dumps(payload, separators=(",", ":")) + "'"


# --------------------------------------------------------------------------- main
def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate + upload an ACT Agent recording to MinIO/S3.")
    parser.add_argument("--py", required=True, help="Path to the recorded Playwright .py")
    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--params", help="Path to a params JSON file")
    src.add_argument("--params-json", help="Inline params JSON string")
    parser.add_argument("--name", help="Recording name (default: .py filename stem)")
    parser.add_argument("--format", choices=["xlsx", "csv"], default="xlsx", help="Params file format (default xlsx)")
    parser.add_argument("--bucket", help="Storage bucket (default: TENANT_ID or STORAGE_ACTIVITIES_BUCKET)")
    parser.add_argument("--env", default=str(Path(__file__).resolve().parent.parent / ".env"), help="dotenv path")
    parser.add_argument("--execution-mode", choices=["parallel", "sequential"], default="parallel")
    parser.add_argument("--dry-run", action="store_true", help="Build files but do not upload")
    args = parser.parse_args(argv)

    py_path = Path(args.py).expanduser()
    if not py_path.is_file():
        raise SystemExit(f"ERROR: --py not found: {py_path}")
    name = args.name or py_path.stem

    if args.params_json is not None:
        payload = json.loads(args.params_json)
    else:
        payload = json.loads(Path(args.params).expanduser().read_text())

    param_sets = normalize_param_sets(payload)

    py_bytes = py_path.read_bytes()
    if args.format == "xlsx":
        params_bytes = build_params_xlsx(param_sets)
        params_ext = "_params.xlsx"
        params_ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        params_bytes = build_params_csv(param_sets)
        params_ext = "_params.csv"
        params_ct = "text/csv"

    folder = f"recordings/{name}/"
    py_key = f"{folder}{name}.py"
    params_key = f"{folder}{name}{params_ext}"

    # Cross-check: every {{placeholder}} in the script is covered by the first param set,
    # so the run won't fail the runner's unresolved-placeholder gate.
    import re as _re
    placeholders = sorted(set(_re.findall(r"\{\{\s*([\w.-]+)\s*\}\}", py_bytes.decode("utf-8", "replace"))))
    missing = [p for p in placeholders if p not in param_sets[0]]

    print(f"recording name : {name}")
    print(f"param sets     : {len(param_sets)} row(s); keys: {', '.join(_ordered_headers(param_sets))}")
    print(f"placeholders   : {', '.join(placeholders) or '(none)'}")
    if missing:
        print(f"WARNING: placeholders with no param value: {', '.join(missing)} -> the run will fail the placeholder gate.")

    if args.dry_run:
        out_dir = py_path.parent
        (out_dir / f"{name}{params_ext}").write_bytes(params_bytes)
        print(f"\n[dry-run] wrote params file locally: {out_dir / (name + params_ext)}")
        print(f"[dry-run] would upload to: <bucket>/{py_key} and <bucket>/{params_key}")
        print("\nrun command:\n" + run_command(name, py_key, args.execution_mode))
        return 0

    cfg = load_env(Path(args.env).expanduser())
    bucket = resolve_bucket(cfg, args.bucket)
    s3 = make_s3(cfg)
    s3.put_object(Bucket=bucket, Key=py_key, Body=py_bytes, ContentType="text/x-python")
    s3.put_object(Bucket=bucket, Key=params_key, Body=params_bytes, ContentType=params_ct)

    # Verify round-trip.
    back_py = s3.get_object(Bucket=bucket, Key=py_key)["Body"].read()
    back_params = s3.get_object(Bucket=bucket, Key=params_key)["Body"].read()
    ok = back_py == py_bytes and back_params == params_bytes
    print(f"\nuploaded to bucket: {bucket}")
    print(f"   {py_key}  ({len(py_bytes)}b)")
    print(f"   {params_key}  ({len(params_bytes)}b)")
    print("verified round-trip: " + ("OK" if ok else "MISMATCH"))
    if not ok:
        return 1

    print("\nrun command (from act-v2/act_agent):\n" + run_command(name, py_key, args.execution_mode))
    return 0


if __name__ == "__main__":
    sys.exit(main())
